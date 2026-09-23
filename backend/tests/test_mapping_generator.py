"""Tests for mapping engine and document generator (Tasks 1.11, 1.12)."""

import io

import pytest
from docx import Document
import openpyxl
from pptx import Presentation
from pptx.util import Inches

from app.services.mapping.parser import parse_template_bytes
from app.services.mapping.mapper import map_fields
from app.services.mapping.generator import generate_filled_document


def make_docx(placeholders: list[str]) -> bytes:
    doc = Document()
    for ph in placeholders:
        doc.add_paragraph(f"Line with {ph} here.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def make_xlsx(cells: dict) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    for coord, val in cells.items():
        ws[coord] = val
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def make_pptx(placeholders: list[str]) -> bytes:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    left, top, width, height = Inches(1), Inches(1), Inches(3), Inches(1)
    shape = slide.shapes.add_textbox(left, top, width, height)
    tf = shape.text_frame
    tf.text = f"Slide {placeholders[0]}"
    if len(placeholders) > 1:
        shape2 = slide.shapes.add_textbox(left, Inches(2), width, height)
        shape2.text_frame.text = f"Second {placeholders[1]}"
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


class TestMapper:
    def test_exact_match(self):
        tpl = make_docx(["{{full_name}}", "{{invoice_date}}"])
        parsed = parse_template_bytes(tpl, filename="t.docx")
        extracted = {"full_name": "John Doe", "invoice_date": "2026-09-23"}
        result = map_fields(extracted, parsed)
        assert len(result.mapped) == 2
        assert result.mapped["{{full_name}}"] == "John Doe"
        assert len(result.unmapped) == 0
        assert all(m.status == "mapped" for m in result.mappings)

    def test_unmapped_field(self):
        tpl = make_docx(["{{phone_number}}", "{{full_name}}"])
        parsed = parse_template_bytes(tpl, filename="t.docx")
        extracted = {"full_name": "Alice"}
        result = map_fields(extracted, parsed)
        assert "phone_number" in result.unmapped
        assert "full_name" not in result.unmapped
        assert len(result.mapped) == 1

    def test_synonym_match(self):
        # extracted uses "nama" (ID) but template has full_name
        tpl = make_docx(["{{full_name}}"])
        parsed = parse_template_bytes(tpl, filename="t.docx")
        extracted = {"nama": "Budi Santoso"}  # synonym for full_name
        result = map_fields(extracted, parsed)
        # Should map via synonym canonicalization
        assert result.mapped.get("{{full_name}}") == "Budi Santoso"
        assert result.mappings[0].status in ("synonym", "mapped")

    def test_fuzzy_match(self):
        tpl = make_docx(["{{total_amount}}"])
        parsed = parse_template_bytes(tpl, filename="t.docx")
        # Intentional typo-ish variant
        extracted = {"total_amout": "$5000"}  # missing n
        result = map_fields(extracted, parsed, fuzzy_cutoff=0.6)
        assert result.mapped.get("{{total_amount}}") == "$5000"
        assert result.mappings[0].status == "fuzzy"

    def test_null_value_treated_as_unmapped(self):
        tpl = make_docx(["{{field1}}"])
        parsed = parse_template_bytes(tpl, filename="t.docx")
        extracted = {"field1": None}
        result = map_fields(extracted, parsed)
        assert "field1" in result.unmapped
        assert len(result.mapped) == 0

    def test_format_preserved(self):
        tpl = make_docx(["{{invoice_date}}"])
        parsed = parse_template_bytes(tpl, filename="t.docx")
        extracted = {"invoice_date": "  2026-09-23  "}
        result = map_fields(extracted, parsed)
        assert result.mapped["{{invoice_date}}"] == "2026-09-23"

    def test_values_by_field_contains_both_keys(self):
        tpl = make_docx(["{{full_name}}"])
        parsed = parse_template_bytes(tpl, filename="t.docx")
        extracted = {"full_name": "John"}
        result = map_fields(extracted, parsed)
        assert result.values_by_field["full_name"] == "John"
        assert result.mapped["{{full_name}}"] == "John"


class TestDocxGenerator:
    def test_fill_single_placeholder(self):
        tpl = make_docx(["{{full_name}}"])
        mapped = {"{{full_name}}": "John Doe"}
        filled = generate_filled_document(tpl, "template.docx", mapped)
        # Verify by reparsing / reading docx
        doc = Document(io.BytesIO(filled))
        text = "\n".join(p.text for p in doc.paragraphs)
        assert "John Doe" in text
        assert "{{full_name}}" not in text

    def test_fill_multiple_placeholders(self):
        tpl = make_docx(["{{name}}", "{{date}}"])
        mapped = {"{{name}}": "Alice", "{{date}}": "2026-09-23"}
        filled = generate_filled_document(tpl, "tpl.docx", mapped)
        doc = Document(io.BytesIO(filled))
        text = "\n".join(p.text for p in doc.paragraphs)
        assert "Alice" in text
        assert "2026-09-23" in text

    def test_fill_via_field_name_key(self):
        tpl = make_docx(["{{full_name}}"])
        # Mapper can give field_name keys instead of raw placeholder
        mapped = {"full_name": "Jane"}
        filled = generate_filled_document(tpl, "t.docx", mapped)
        doc = Document(io.BytesIO(filled))
        text = "\n".join(p.text for p in doc.paragraphs)
        assert "Jane" in text

    def test_leave_unmapped_placeholder(self):
        tpl = make_docx(["{{found}}", "{{not_found}}"])
        mapped = {"{{found}}": "value"}
        filled = generate_filled_document(tpl, "t.docx", mapped)
        doc = Document(io.BytesIO(filled))
        text = "\n".join(p.text for p in doc.paragraphs)
        assert "value" in text
        assert "{{not_found}}" in text  # preserved

    def test_preserve_table(self):
        doc = Document()
        table = doc.add_table(rows=1, cols=2)
        table.rows[0].cells[0].text = "{{field1}}"
        table.rows[0].cells[1].text = "static"
        buf = io.BytesIO()
        doc.save(buf)
        filled = generate_filled_document(buf.getvalue(), "t.docx", {"{{field1}}": "replaced"})
        doc2 = Document(io.BytesIO(filled))
        assert doc2.tables[0].rows[0].cells[0].text == "replaced"

    def test_mixed_pattern_types(self):
        tpl = make_docx(["{{a}}", "[b]", "<<c>>", "__d__"])
        mapped = {"{{a}}": "1", "[b]": "2", "<<c>>": "3", "__d__": "4"}
        filled = generate_filled_document(tpl, "t.docx", mapped)
        doc = Document(io.BytesIO(filled))
        text = "\n".join(p.text for p in doc.paragraphs)
        assert "1" in text and "2" in text and "3" in text and "4" in text

    def test_empty_mapped_returns_original(self):
        tpl = make_docx(["{{field}}"])
        filled = generate_filled_document(tpl, "t.docx", {})
        doc = Document(io.BytesIO(filled))
        text = "\n".join(p.text for p in doc.paragraphs)
        assert "{{field}}" in text


class TestXlsxGenerator:
    def test_fill_cell(self):
        tpl = make_xlsx({"A1": "{{name}}", "B2": "Total: <<amount>>"})
        mapped = {"{{name}}": "Bob", "<<amount>>": "$100"}
        filled = generate_filled_document(tpl, "t.xlsx", mapped)
        wb = openpyxl.load_workbook(io.BytesIO(filled))
        ws = wb.active
        assert ws["A1"].value == "Bob"
        assert ws["B2"].value == "Total: $100"

    def test_leave_unmapped(self):
        tpl = make_xlsx({"A1": "{{a}}", "A2": "{{b}}"})
        mapped = {"{{a}}": "1"}
        filled = generate_filled_document(tpl, "t.xlsx", mapped)
        wb = openpyxl.load_workbook(io.BytesIO(filled))
        ws = wb.active
        assert ws["A1"].value == "1"
        assert ws["A2"].value == "{{b}}"


class TestPptxGenerator:
    def test_fill_textbox(self):
        tpl = make_pptx(["{{title}}"])
        mapped = {"{{title}}": "New Title"}
        filled = generate_filled_document(tpl, "t.pptx", mapped)
        prs = Presentation(io.BytesIO(filled))
        texts = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if shape.has_text_frame:
                    texts.append(shape.text_frame.text)
        assert any("New Title" in t for t in texts)
        assert not any("{{title}}" in t for t in texts)

    def test_preserve_unmapped(self):
        tpl = make_pptx(["{{a}}", "{{b}}"])
        mapped = {"{{a}}": "1"}
        filled = generate_filled_document(tpl, "t.pptx", mapped)
        prs = Presentation(io.BytesIO(filled))
        texts = " ".join(shape.text_frame.text for slide in prs.slides for shape in slide.shapes if shape.has_text_frame)
        assert "1" in texts
        assert "{{b}}" in texts

    def test_unsupported_format_raises(self):
        with pytest.raises(ValueError, match="Unsupported"):
            generate_filled_document(b"fake", "template.pdf", {})

    def test_empty_template_raises(self):
        with pytest.raises(ValueError, match="Empty"):
            generate_filled_document(b"", "t.docx", {})


class TestEndToEndMappingGeneration:
    def test_full_flow_docx(self):
        # Parse -> map -> generate -> verify
        tpl = make_docx(["Hello {{full_name}}, your invoice {{invoice_number}} total is <<total_amount>>."])
        parsed = parse_template_bytes(tpl, filename="inv.docx")
        assert len(parsed.fields) == 3
        extracted = {"full_name": "John Doe", "invoice_number": "INV-001", "total_amount": "$5,000"}
        mapped = map_fields(extracted, parsed)
        assert len(mapped.unmapped) == 0
        filled = generate_filled_document(tpl, "inv.docx", mapped.mapped)
        doc = Document(io.BytesIO(filled))
        text = "\n".join(p.text for p in doc.paragraphs)
        assert "John Doe" in text
        assert "INV-001" in text
        assert "$5,000" in text
        assert "{{" not in text
