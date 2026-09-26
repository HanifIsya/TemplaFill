"""Security hardening regression tests (VULN-01..VULN-15).

Covers the fixes made to address the cyber security assessment.
"""

import asyncio
import io
import time

import openpyxl
import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from starlette.requests import Request
from starlette.responses import PlainTextResponse

from app.api.upload import _zip_bomb_guard
from app.core.security import (
    InMemoryRateLimiter,
    SizeLimitedMiddleware,
    create_session_token,
    get_client_ip,
    verify_session_token,
)
from app.services.mapping.generator import neutralize_formula


# ---------------------------------------------------------------------------
# VULN-01: session tokens
# ---------------------------------------------------------------------------

class TestSessionTokens:
    def test_valid_token_verifies(self):
        token = create_session_token("job-123")
        assert verify_session_token(token, "job-123") is True

    def test_token_for_other_job_rejected(self):
        token = create_session_token("job-123")
        assert verify_session_token(token, "job-999") is False

    def test_tampered_token_rejected(self):
        token = create_session_token("job-123")
        tampered = token[:-2] + ("aa" if not token.endswith("aa") else "bb")
        assert verify_session_token(tampered, "job-123") is False

    def test_garbage_token_rejected(self):
        assert verify_session_token("not-a-token", "job-123") is False
        assert verify_session_token("", "job-123") is False

    def test_bearer_prefix_accepted(self):
        token = create_session_token("job-123")
        assert verify_session_token(f"Bearer {token}", "job-123") is True


# ---------------------------------------------------------------------------
# VULN-04: formula / DDE neutralization
# ---------------------------------------------------------------------------

class TestFormulaNeutralization:
    @pytest.mark.parametrize(
        "payload",
        ["=cmd|'/C calc'!A1", '=HYPERLINK("http://evil")', "@SUM(A1)", "+cmd|'/C calc'!A1", "=1+1"],
    )
    def test_dangerous_values_neutralized(self, payload):
        out = neutralize_formula(payload)
        # The apostrophe is Excel's literal-text escape; the original value is
        # preserved after it so it is stored (not evaluated) as text.
        assert out == "'" + payload, f"{payload!r} was not neutralized"

    @pytest.mark.parametrize("safe", ["-500", "+12.5", "1234", "John Doe", "-1,250.00", "$500"])
    def test_legitimate_values_preserved(self, safe):
        assert neutralize_formula(safe) == safe

    def test_leading_tab_neutralized(self):
        assert neutralize_formula("\t=1+1") == "'\t=1+1"

    def test_xlsx_generation_neutralizes_formula(self):
        from app.services.mapping.generator import generate_filled_document

        wb = openpyxl.Workbook()
        ws = wb.active
        ws["A1"] = "Total: {{amount}}"
        buf = io.BytesIO()
        wb.save(buf)

        filled = generate_filled_document(buf.getvalue(), "t.xlsx", {"{{amount}}": "=cmd|'/C calc'!A1"})
        out = openpyxl.load_workbook(io.BytesIO(filled)).active["A1"].value
        # The dangerous payload is stored as literal text, not a live formula.
        assert out == "Total: '=cmd|'/C calc'!A1"


# ---------------------------------------------------------------------------
# VULN-05: rate limiter bounded + trusted proxy IP
# ---------------------------------------------------------------------------

class TestRateLimiter:
    def test_limit_enforced(self):
        limiter = InMemoryRateLimiter()
        for _ in range(60):
            assert limiter.is_allowed("1.2.3.4", "read") is True
        assert limiter.is_allowed("1.2.3.4", "read") is False

    def test_buckets_bounded(self):
        limiter = InMemoryRateLimiter()
        limiter.MAX_TRACKED_KEYS = 50
        for i in range(500):
            limiter.is_allowed(f"10.0.{i // 256}.{i % 256}", "read")
        # Prune empty buckets; never creep beyond the cap once buckets drain.
        for key in list(limiter._buckets.keys()):
            limiter._buckets[key].clear()
        limiter.is_allowed("another", "read")
        assert len(limiter._buckets) <= limiter.MAX_TRACKED_KEYS


class TestClientIp:
    def test_ignores_xff_from_untrusted_peer(self):
        scope = {
            "type": "http",
            "headers": [(b"x-forwarded-for", b"9.9.9.9")],
            "client": ("203.0.113.5", 1234),
            "method": "GET",
            "path": "/",
        }
        req = Request(scope)
        assert get_client_ip(req) == "203.0.113.5"


# ---------------------------------------------------------------------------
# VULN-11: ZIP-bomb guard
# ---------------------------------------------------------------------------

class TestZipBombGuard:
    def test_normal_template_passes(self):
        wb = openpyxl.Workbook()
        wb.active["A1"] = "x"
        buf = io.BytesIO()
        wb.save(buf)
        assert _zip_bomb_guard(buf.getvalue()) is None

    def test_high_ratio_archive_rejected(self):
        import zipfile

        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
            zf.writestr("word/document.xml", "A" * (5 * 1024 * 1024))
        assert _zip_bomb_guard(buf.getvalue()) is not None

    def test_invalid_zip_rejected(self):
        assert _zip_bomb_guard(b"PK\x03\x04not a real zip") is not None


# ---------------------------------------------------------------------------
# VULN-03: request body size middleware
# ---------------------------------------------------------------------------

class TestSizeLimitMiddleware:
    def _app(self, max_bytes: int) -> TestClient:
        app = FastAPI()

        @app.post("/echo")
        async def echo():
            return PlainTextResponse("ok")

        app.add_middleware(SizeLimitedMiddleware, max_body_bytes=max_bytes)
        return TestClient(app)

    def test_oversized_request_rejected(self):
        client = self._app(10)
        resp = client.post("/echo", content=b"x" * 100)
        assert resp.status_code == 413
        assert resp.json()["error"]["code"] == "FILE_TOO_LARGE"

    def test_small_request_allowed(self):
        client = self._app(1000)
        assert client.post("/echo", content=b"x" * 10).status_code == 200


# ---------------------------------------------------------------------------
# VULN-02: job eviction / capacity
# ---------------------------------------------------------------------------

class TestJobEviction:
    def test_capacity_enforced(self):
        from app.services.jobs.manager import JobManager, _jobs

        manager = JobManager()
        _jobs.clear()

        async def run():
            for i in range(12):
                await manager.create_job(
                    source_bytes=b"%PDF-1.4 test",
                    source_filename=f"s{i}.pdf",
                    template_bytes=b"PK\x03\x04",
                    template_filename=f"t{i}.docx",
                )
            manager._evict_locked(retention_hours=24, max_jobs=5)
            return len(_jobs)

        remaining = asyncio.run(run())
        _jobs.clear()
        assert remaining == 5

    def test_expired_jobs_evicted(self):
        from app.services.jobs.manager import JobManager, _jobs

        manager = JobManager()
        _jobs.clear()

        async def run():
            await manager.create_job(
                source_bytes=b"%PDF-1.4 test",
                source_filename="s.pdf",
                template_bytes=b"PK\x03\x04",
                template_filename="t.docx",
            )
            for job in _jobs.values():
                job.created_at_monotonic = time.monotonic() - 999999
            manager._evict_locked(retention_hours=1, max_jobs=100)
            return list(_jobs.values())

        remaining = asyncio.run(run())
        _jobs.clear()
        assert remaining == []


# ---------------------------------------------------------------------------
# VULN-14: production secret enforcement
# ---------------------------------------------------------------------------

class TestProductionSecret:
    def test_placeholder_secret_rejected_in_production(self):
        from app.core.config import Settings

        with pytest.raises(Exception):
            Settings(app_env="production", secret_key="change_this_to_a_random_string_in_production")


# ---------------------------------------------------------------------------
# VULN-06: Gemini key not in URL
# ---------------------------------------------------------------------------

class TestGeminiRestKeyHandling:
    def test_key_not_in_url(self):
        import inspect

        from app.services.generation.extractor import GeminiExtractor

        src = inspect.getsource(GeminiExtractor._call_gemini_rest)
        assert "?key=" not in src
        assert "x-goog-api-key" in src


# ---------------------------------------------------------------------------
# Sanitized routing
# ---------------------------------------------------------------------------

class TestSanitizedErrors:
    def test_unknown_route_404(self):
        from app.main import app

        client = TestClient(app)
        assert client.get("/api/definitely-not-a-route").status_code == 404