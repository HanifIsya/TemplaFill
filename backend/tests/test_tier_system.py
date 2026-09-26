"""Tier system tests — Phase 6 (ADR-020): auth, quotas, provider selection.

Covers TESTING.md tier matrix T1–T8 (backend rows) and TIER_ARCHITECTURE §8:
 - tier/job token purpose isolation
 - scrypt credential verification
 - login/logout/quota endpoints + auth rate limit
 - daily free/pro quotas (429 QUOTA_EXCEEDED + Retry-After)
 - provider selection by Job.tier
 - DeepSeek extractor: fake mode, request building, heuristic-only fallback
 - zero-Gemini guarantee on the pro path (Google calls monkeypatched to raise)
"""

import hashlib
import hmac
import io
import json
import time
import uuid

import fitz  # PyMuPDF
import pytest
from fastapi.testclient import TestClient

from app.core.config import get_settings
from app.core.security import (
    _b64url_encode,
    _signing_key,
    create_session_token,
    create_tier_token,
    get_rate_limiter,
    verify_session_token,
    verify_tier_token,
)
from app.core.tier_auth import hash_password, verify_password
from app.main import app
from app.services.jobs.manager import _jobs, get_job_manager

client = TestClient(app)
manager = get_job_manager()


# ---------------------------------------------------------------------------
# Helpers & fixtures
# ---------------------------------------------------------------------------

def create_pdf_bytes(text: str, title: str = "Test PDF") -> bytes:
    doc = fitz.open()
    page = doc.new_page()
    page.insert_textbox(fitz.Rect(50, 50, 550, 800), text, fontsize=11, fontname="helv")
    doc.set_metadata({"title": title})
    buf = io.BytesIO()
    doc.save(buf)
    doc.close()
    return buf.getvalue()


def create_docx_bytes(placeholders: list[str]) -> bytes:
    from docx import Document

    doc = Document()
    for ph in placeholders:
        doc.add_paragraph(f"Field {ph}")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def upload_files():
    pdf = create_pdf_bytes("Applicant: John Doe\nInvoice: INV-001\nEmail: john@example.com")
    docx = create_docx_bytes(["{{full_name}}", "{{invoice_number}}"])
    return {
        "source_file": ("source.pdf", pdf, "application/pdf"),
        "template_file": ("template.docx", docx, "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
    }


@pytest.fixture(autouse=True)
def clean_state():
    """Isolate jobs + rate-limiter buckets + cookies per test.

    The module-level TestClient shares a cookie jar; without clearing it a
    login test would leak the tier cookie into later upload tests.
    """
    _jobs.clear()
    get_rate_limiter().clear()
    client.cookies.clear()
    yield
    _jobs.clear()
    get_rate_limiter().clear()
    client.cookies.clear()


@pytest.fixture()
def patch_client_ip(monkeypatch):
    """Route HTTP rate/quota checks to a fresh synthetic IP.

    TestClient hard-codes the peer host as 'testclient' (which bypasses rate
    limits), so we override the IP resolver inside the dependencies module.
    """
    used = {"ip": f"203.0.113.{uuid.uuid4().int % 200 + 1}"}

    def _fake_ip(request):
        return used["ip"]

    monkeypatch.setattr("app.api.dependencies.get_client_ip", _fake_ip)
    return used


@pytest.fixture()
def configured_account(monkeypatch):
    settings = get_settings()
    monkeypatch.setattr(settings, "tier_account_username", "tester")
    monkeypatch.setattr(settings, "tier_account_password_hash", hash_password("s3cret-pass"))
    return settings


def _boom(*_args, **_kwargs):
    raise AssertionError("Google/Gemini must never be called on the account-tier path")


def _expired_tier_token() -> str:
    payload = f"tier.pro.{int(time.time()) - 60}.nonce".encode("utf-8")
    encoded = _b64url_encode(payload)
    sig = hmac.new(_signing_key(), encoded.encode("ascii"), hashlib.sha256).digest()
    return f"{encoded}.{_b64url_encode(sig)}"


# ---------------------------------------------------------------------------
# Tier tokens — purpose isolation (ADR-020 §4)
# ---------------------------------------------------------------------------

class TestTierTokens:
    def test_tier_token_roundtrip(self):
        token = create_tier_token(tier="pro")
        assert verify_tier_token(token) == "pro"

    def test_job_token_is_not_a_tier_token(self):
        job_token = create_session_token("11111111-1111-1111-1111-111111111111")
        assert verify_tier_token(job_token) is None

    def test_tier_token_is_not_a_job_token(self):
        tier_token = create_tier_token(tier="pro")
        assert verify_session_token(tier_token, "11111111-1111-1111-1111-111111111111") is False

    def test_expired_tier_token_rejected(self):
        assert verify_tier_token(_expired_tier_token()) is None

    def test_tampered_tier_token_rejected(self):
        token = create_tier_token(tier="pro")
        # Flip a signature character well inside the base64 data. (The LAST
        # character is a bad choice: its unused bits are ignored by the b64
        # decoder, so a flip there can still decode to the same bytes.)
        dot = token.index(".")
        idx = dot + 10
        flipped = "y" if token[idx] != "y" else "x"
        assert verify_tier_token(token[:idx] + flipped + token[idx + 1 :]) is None

    def test_empty_and_garbage_tokens_rejected(self):
        assert verify_tier_token("") is None
        assert verify_tier_token("not-a-token") is None
        assert verify_tier_token(None) is None  # type: ignore[arg-type]


# ---------------------------------------------------------------------------
# scrypt shared credential (ADR-020 §4)
# ---------------------------------------------------------------------------

class TestSharedPassword:
    def test_roundtrip(self):
        stored = hash_password("hunter2")
        assert stored.startswith("scrypt$")
        assert verify_password("hunter2", stored) is True

    def test_wrong_password_rejected(self):
        stored = hash_password("hunter2")
        assert verify_password("wrong", stored) is False

    def test_malformed_hash_fails_closed(self):
        assert verify_password("x", "") is False
        assert verify_password("x", "not-a-hash") is False
        assert verify_password("x", "bcrypt$aa$bb") is False
        assert verify_password("x", "scrypt$zz$yy") is False  # bad hex

    def test_fresh_salt_each_hash(self):
        assert hash_password("same") != hash_password("same")

    def test_hash_script_importable(self):
        import pathlib
        import subprocess
        import sys

        script = pathlib.Path(__file__).resolve().parents[2] / "scripts" / "hash_password.py"
        out = subprocess.run(
            [sys.executable, str(script), "abc123"],
            capture_output=True,
            text=True,
            cwd=str(pathlib.Path(__file__).resolve().parents[2]),
            timeout=30,
        )
        assert out.returncode == 0, out.stderr
        assert verify_password("abc123", out.stdout.strip()) is True


# ---------------------------------------------------------------------------
# Login / logout / quota endpoints
# ---------------------------------------------------------------------------

class TestAuthEndpoints:
    def test_login_disabled_without_password_hash(self, monkeypatch):
        monkeypatch.setattr(get_settings(), "tier_account_password_hash", "")
        resp = client.post("/api/auth/login", json={"username": "a", "password": "b"})
        assert resp.status_code == 503
        assert resp.json()["error"]["code"] == "TIER_NOT_CONFIGURED"

    def test_login_success(self, configured_account):
        resp = client.post("/api/auth/login", json={"username": "tester", "password": "s3cret-pass"})
        assert resp.status_code == 200, resp.text
        data = resp.json()["data"]
        assert data["tier"] == "pro"
        assert verify_tier_token(data["token"]) == "pro"
        assert configured_account.tier_cookie_name in resp.cookies

    def test_login_wrong_password_generic_401(self, configured_account):
        resp = client.post("/api/auth/login", json={"username": "tester", "password": "nope"})
        assert resp.status_code == 401
        assert resp.json()["error"]["code"] == "INVALID_CREDENTIALS"

    def test_login_wrong_username_same_message_as_password(self, configured_account):
        resp = client.post("/api/auth/login", json={"username": "ghost", "password": "s3cret-pass"})
        assert resp.status_code == 401
        assert resp.json()["error"]["code"] == "INVALID_CREDENTIALS"

    def test_auth_rate_limit_5_per_minute(self, configured_account, patch_client_ip):
        url = "/api/auth/login"
        for i in range(5):
            resp = client.post(url, json={"username": "tester", "password": "wrong"})
            assert resp.status_code == 401, f"attempt {i}: {resp.status_code}"
        resp = client.post(url, json={"username": "tester", "password": "wrong"})
        assert resp.status_code == 429
        assert resp.json()["error"]["code"] == "RATE_LIMITED"
        assert "Retry-After" in resp.headers

    def test_logout_clears_tier_cookie(self, configured_account):
        login = client.post("/api/auth/login", json={"username": "tester", "password": "s3cret-pass"})
        assert login.status_code == 200
        resp = client.post("/api/auth/logout")
        assert resp.status_code == 200
        assert resp.json()["data"]["tier"] == "free"

    def test_quota_endpoint_reports_tier_and_limits(self, configured_account):
        resp = client.get("/api/auth/quota")
        assert resp.status_code == 200
        data = resp.json()["data"]
        assert data["tier"] == "free"
        assert data["free_limit"] == get_settings().free_jobs_per_day
        assert data["pro_limit"] == get_settings().pro_jobs_per_day
        assert data["free_used_today"] == 0

        resp = client.get("/api/auth/quota", headers={"X-Tier-Token": create_tier_token()})
        assert resp.json()["data"]["tier"] == "pro"


# ---------------------------------------------------------------------------
# Daily quotas (ADR-020 §5)
# ---------------------------------------------------------------------------

class TestDailyQuota:
    def test_limiter_day_window_enforces_free_cap(self):
        settings = get_settings()
        limiter = get_rate_limiter()
        ip = "198.51.100.10"
        for i in range(settings.free_jobs_per_day):
            assert limiter.is_allowed(ip, "free_upload") is True, f"slot {i + 1}"
        assert limiter.is_allowed(ip, "free_upload") is False
        assert limiter.usage(ip, "free_upload") == settings.free_jobs_per_day
        assert limiter.retry_after(ip, "free_upload") > 0

    def test_pro_window_independent_of_free(self):
        settings = get_settings()
        limiter = get_rate_limiter()
        ip = "198.51.100.11"
        for _ in range(settings.free_jobs_per_day):
            assert limiter.is_allowed(ip, "free_upload") is True
        assert limiter.is_allowed(ip, "free_upload") is False
        # Pro bucket untouched by free consumption
        assert limiter.usage(ip, "pro_upload") == 0
        assert limiter.is_allowed(ip, "pro_upload") is True

    def test_free_uploads_capped_at_5_with_quota_exceeded(self, patch_client_ip):
        ip = patch_client_ip["ip"]
        for i in range(5):
            resp = client.post("/api/upload", files=upload_files())
            assert resp.status_code == 202, f"upload {i + 1}: {resp.status_code} {resp.text}"
        resp = client.post("/api/upload", files=upload_files())
        assert resp.status_code == 429
        body = resp.json()
        assert body["error"]["code"] == "QUOTA_EXCEEDED"
        assert "Retry-After" in resp.headers
        assert body["error"]["details"]["tier"] == "free"
        assert get_rate_limiter().usage(ip, "free_upload") == 5

    def test_pro_upload_consumes_pro_bucket_only(self, patch_client_ip, configured_account):
        ip = patch_client_ip["ip"]
        resp = client.post(
            "/api/upload",
            files=upload_files(),
            headers={"X-Tier-Token": create_tier_token()},
        )
        assert resp.status_code == 202, resp.text
        assert resp.json()["data"]["tier"] == "pro"
        limiter = get_rate_limiter()
        assert limiter.usage(ip, "pro_upload") == 1
        assert limiter.usage(ip, "free_upload") == 0


# ---------------------------------------------------------------------------
# Provider selection & Job.tier binding
# ---------------------------------------------------------------------------

class TestProviderSelection:
    def test_tier_factories(self):
        from app.services.generation.deepseek_extractor import DeepSeekExtractor
        from app.services.generation.extractor import GeminiExtractor, get_extractor_for_tier

        free = get_extractor_for_tier("free", force_fake=True)
        pro = get_extractor_for_tier("pro", force_fake=True)
        assert isinstance(free, GeminiExtractor)
        assert isinstance(pro, DeepSeekExtractor)

    def test_create_job_defaults_to_free(self):
        import asyncio

        job = asyncio.run(
            manager.create_job(
                source_bytes=b"%PDF-1.4 x",
                source_filename="s.pdf",
                template_bytes=b"PK\x03\x04",
                template_filename="t.docx",
            )
        )
        assert job.tier == "free"

    def test_create_job_pro_tier(self):
        import asyncio

        job = asyncio.run(
            manager.create_job(
                source_bytes=b"%PDF-1.4 x",
                source_filename="s.pdf",
                template_bytes=b"PK\x03\x04",
                template_filename="t.docx",
                tier="pro",
            )
        )
        assert job.tier == "pro"

    def test_upload_without_token_is_free(self):
        resp = client.post("/api/upload", files=upload_files())
        assert resp.status_code == 202
        assert resp.json()["data"]["tier"] == "free"

    def test_upload_with_valid_tier_token_is_pro(self):
        resp = client.post(
            "/api/upload",
            files=upload_files(),
            headers={"X-Tier-Token": create_tier_token()},
        )
        assert resp.status_code == 202
        assert resp.json()["data"]["tier"] == "pro"

    def test_upload_with_garbage_tier_token_stays_free(self):
        resp = client.post(
            "/api/upload",
            files=upload_files(),
            headers={"X-Tier-Token": "bogus.token.value"},
        )
        assert resp.status_code == 202
        assert resp.json()["data"]["tier"] == "free"

    def test_upload_with_expired_tier_token_stays_free(self):
        resp = client.post(
            "/api/upload",
            files=upload_files(),
            headers={"X-Tier-Token": _expired_tier_token()},
        )
        assert resp.status_code == 202
        assert resp.json()["data"]["tier"] == "free"


# ---------------------------------------------------------------------------
# DeepSeekExtractor
# ---------------------------------------------------------------------------

class TestDeepSeekExtractor:
    async def test_fake_mode_uses_heuristic_only(self):
        from app.services.generation.deepseek_extractor import get_deepseek_extractor

        ext = get_deepseek_extractor(force_fake=True)
        res = await ext.extract_batch(
            [{"field_name": "full_name", "description": "Name"}],
            ["Applicant: John Doe"],
            source_pages=[1],
        )
        assert res["full_name"].extracted_by == "heuristic"
        assert ext.last_engine_used == "heuristic"

    async def test_request_building_and_parse(self, monkeypatch):
        from app.services.generation.deepseek_extractor import DeepSeekExtractor

        captured = {}

        async def fake_call(self, prompt):
            captured["prompt"] = prompt
            return json.dumps(
                {
                    "extractions": [
                        {
                            "field_name": "full_name",
                            "value": "John Doe",
                            "confidence": 0.99,
                            "source_page": 1,
                            "source_text": "Applicant: John Doe",
                        }
                    ]
                }
            )

        monkeypatch.setattr(DeepSeekExtractor, "_call_deepseek", fake_call)
        ext = DeepSeekExtractor(api_key="sk-test", use_fake=False)
        res = await ext.extract_batch(
            [{"field_name": "full_name", "description": "Full name"}],
            ["Applicant: John Doe"],
            source_pages=[7],
        )
        # Shared prompt builder with untrusted-data fencing (VULN-07)
        assert "<document>" in captured["prompt"]
        assert "full_name" in captured["prompt"]
        assert res["full_name"].extracted_by == "deepseek"
        assert res["full_name"].extracted_value == "John Doe"
        assert res["full_name"].source_page == 7  # remapped through source_pages
        assert ext.last_engine_used == "deepseek"

    async def test_missing_field_recovered_by_heuristic(self, monkeypatch):
        from app.services.generation.deepseek_extractor import DeepSeekExtractor

        async def fake_call(self, prompt):
            return json.dumps({"extractions": []})

        monkeypatch.setattr(DeepSeekExtractor, "_call_deepseek", fake_call)
        ext = DeepSeekExtractor(api_key="sk-test", use_fake=False)
        res = await ext.extract_batch(
            [{"field_name": "full_name", "description": "Full name"}],
            ["Applicant: John Doe"],
        )
        assert res["full_name"].extracted_by == "heuristic"
        assert ext.last_engine_used == "hybrid"

    async def test_api_failure_falls_back_to_heuristic_never_gemini(self, monkeypatch):
        from app.services.generation.deepseek_extractor import DeepSeekExtractor
        from app.services.generation.extractor import GeminiExtractor

        async def failing_call(self, prompt):
            raise RuntimeError("HTTP 500")

        monkeypatch.setattr(DeepSeekExtractor, "_call_deepseek", failing_call)
        # If anything in this path touches Gemini, the test fails loudly.
        monkeypatch.setattr(GeminiExtractor, "extract_batch", _boom)
        monkeypatch.setattr(GeminiExtractor, "extract", _boom)

        ext = DeepSeekExtractor(api_key="sk-test", use_fake=False)
        res = await ext.extract_batch(
            [{"field_name": "full_name", "description": "Full name"}],
            ["Applicant: John Doe"],
        )
        assert res["full_name"].extracted_by == "heuristic"
        assert ext.last_engine_used == "heuristic"
        assert "DeepSeek" in (ext.last_fallback_reason or "")

    async def test_single_field_extract_delegates_to_batch(self, monkeypatch):
        from app.services.generation.deepseek_extractor import DeepSeekExtractor

        async def fake_call(self, prompt):
            return json.dumps(
                {
                    "extractions": [
                        {"field_name": "invoice_number", "value": "INV-001", "confidence": 0.9, "source_page": 1, "source_text": "Invoice: INV-001"}
                    ]
                }
            )

        monkeypatch.setattr(DeepSeekExtractor, "_call_deepseek", fake_call)
        ext = DeepSeekExtractor(api_key="sk-test", use_fake=False)
        res = await ext.extract("invoice_number", ["Invoice: INV-001"])
        assert res.extracted_value == "INV-001"
        assert res.extracted_by == "deepseek"

    async def test_prompt_injected_field_names_ignored(self, monkeypatch):
        """VULN-07: response fields not requested are dropped."""
        from app.services.generation.deepseek_extractor import DeepSeekExtractor

        async def fake_call(self, prompt):
            return json.dumps(
                {
                    "extractions": [
                        {"field_name": "full_name", "value": "John Doe", "confidence": 0.9, "source_page": 1, "source_text": "John Doe"},
                        {"field_name": "evil_field", "value": "injected", "confidence": 1.0, "source_page": 1, "source_text": "x"},
                    ]
                }
            )

        monkeypatch.setattr(DeepSeekExtractor, "_call_deepseek", fake_call)
        ext = DeepSeekExtractor(api_key="sk-test", use_fake=False)
        res = await ext.extract_batch(
            [{"field_name": "full_name", "description": "Full name"}],
            ["Applicant: John Doe"],
        )
        assert "evil_field" not in res
        assert "full_name" in res


# ---------------------------------------------------------------------------
# Sequential batching helper (ADR-020 §3)
# ---------------------------------------------------------------------------

class TestSequentialBatching:
    def test_split_by_chars_preserves_order(self):
        from app.services.jobs.manager import _split_chunks_by_chars

        class C:
            def __init__(self, text):
                self.text = text

        chunks = [C("a" * 100), C("b" * 100), C("c" * 100)]
        batches = _split_chunks_by_chars(chunks, 250)
        assert [[c.text[0] for c in b] for b in batches] == [["a", "b"], ["c"]]

    def test_single_batch_when_within_budget(self):
        from app.services.jobs.manager import _split_chunks_by_chars

        class C:
            def __init__(self, text):
                self.text = text

        chunks = [C("a" * 10), C("b" * 10)]
        batches = _split_chunks_by_chars(chunks, 10_000)
        assert len(batches) == 1
        assert batches[0] == chunks


# ---------------------------------------------------------------------------
# Zero-Gemini guarantee on the pro pipeline (TIER_ARCHITECTURE §1 hard rule)
# ---------------------------------------------------------------------------

class TestProPipelineZeroGoogle:
    def test_pro_job_completes_without_any_google_call(self, monkeypatch):
        import asyncio

        from app.services.generation.extractor import GeminiExtractor
        import app.services.rag.embedder as embedder_mod

        # Any Google touchpoint explodes the test.
        monkeypatch.setattr(GeminiExtractor, "extract_batch", _boom)
        monkeypatch.setattr(GeminiExtractor, "extract", _boom)
        monkeypatch.setattr(embedder_mod, "get_embedder", _boom)

        job = asyncio.run(
            manager.create_job(
                source_bytes=create_pdf_bytes("Applicant: John Doe\nInvoice: INV-001"),
                source_filename="source.pdf",
                template_bytes=create_docx_bytes(["{{full_name}}", "{{invoice_number}}"]),
                template_filename="template.docx",
                tier="pro",
            )
        )
        asyncio.run(manager.process_job(job.job_id))

        assert job.status.value == "completed", job.error
        # RAG bypassed: no vector store, no embeddings
        assert getattr(job, "_vector_store", None) is None
        assert getattr(job, "_chunks", None)
        # Provenance: DeepSeek path (fake mode -> heuristic), never gemini
        assert job.field_results
        for fr in job.field_results:
            assert fr.extracted_by in ("deepseek", "heuristic"), fr.extracted_by
        results = job.to_results_dict()
        assert results["engine_used"] in ("deepseek", "hybrid", "heuristic")
        assert results["tier"] == "pro"

    def test_free_job_still_uses_gemini_path(self):
        """Sanity: the free path is untouched (engine falls back to heuristic offline)."""
        import asyncio

        job = asyncio.run(
            manager.create_job(
                source_bytes=create_pdf_bytes("Applicant: John Doe"),
                source_filename="source.pdf",
                template_bytes=create_docx_bytes(["{{full_name}}"]),
                template_filename="template.docx",
            )
        )
        asyncio.run(manager.process_job(job.job_id))
        assert job.status.value == "completed", job.error
        assert job.tier == "free"
        assert getattr(job, "_vector_store", None) is not None  # RAG active on free tier

    def test_job_status_dict_exposes_tier(self):
        from app.services.jobs.models import Job

        job = Job(tier="pro")
        assert job.to_status_dict()["tier"] == "pro"
        assert job.to_results_dict()["tier"] == "pro"
