"""Regression tests for the backend review fixes (2026-09-26).

Each test pins a specific finding from the full-backend review so the bug
cannot silently return. Grouped by module.
"""

import io

import openpyxl
import pytest
from docx import Document

from app.core.key_pool import GeminiKeyPool, get_key_pool
from app.services.extraction.table_extractor import _table_to_extracted
from app.services.mapping.generator import (
    _build_replacement_map,
    _clean_field_value,
    _replace_in_text,
    generate_filled_document,
)
from app.services.mapping.mapper import map_fields
from app.services.mapping.parser import parse_template_bytes
from app.services.privacy.masker import SelectivePIIMasker
from app.services.rag.chunker import _apply_overlap, chunk_document
from app.services.rag.embedder import Embedder, _fake_embedding
from app.services.rag.vector_store import InMemoryVectorStore, StoredChunk


def _docx_bytes(build) -> bytes:
    doc = Document()
    build(doc)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Generator: double-fill / corruption
# ---------------------------------------------------------------------------

class TestGeneratorDoubleFill:
    def test_value_containing_placeholder_is_not_rescanned(self):
        """A substituted value that itself contains {{...}} must stay verbatim."""
        mapped = {"{{greeting}}": "Hello {{name}}", "{{name}}": "Alice"}
        direct, norm = _build_replacement_map(mapped)
        assert _replace_in_text("{{greeting}}", direct, norm) == "Hello {{name}}"

    def test_nested_delimiter_replaced_once(self):
        # `<<<a>>>` contains `<<a>>`; the longest enclosing match wins so no
        # wrapper residue is left behind.
        direct, norm = _build_replacement_map({"<<<a>>>": "V"})
        assert _replace_in_text("<<<a>>>", direct, norm) == "V"

    def test_docx_value_with_placeholder_survives(self):
        tpl = _docx_bytes(lambda d: d.add_paragraph("{{greeting}}"))
        filled = generate_filled_document(tpl, "t.docx", {"{{greeting}}": "Hello {{name}}"})
        text = "\n".join(p.text for p in Document(io.BytesIO(filled)).paragraphs)
        assert text == "Hello {{name}}"


class TestGeneratorXlsxFormulas:
    def test_existing_formula_cells_are_never_rewritten(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws["A1"] = "=SUM(Table1[Amount])"
        ws["A2"] = "{{name}}"
        buf = io.BytesIO()
        wb.save(buf)

        filled = generate_filled_document(buf.getvalue(), "t.xlsx", {"{{name}}": "Bob"})
        out = openpyxl.load_workbook(io.BytesIO(filled)).active
        assert out["A1"].value == "=SUM(Table1[Amount])"
        assert out["A2"].value == "Bob"

    def test_parser_ignores_structured_references(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws["A1"] = "=SUM(Table1[Amount])"
        ws["A2"] = "{{name}}"
        buf = io.BytesIO()
        wb.save(buf)
        parsed = parse_template_bytes(buf.getvalue(), filename="t.xlsx")
        assert "amount" not in parsed.field_names
        assert "name" in parsed.field_names

    def test_formula_context_template_stored_as_literal(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws["A1"] = "={{amount}}"
        buf = io.BytesIO()
        wb.save(buf)

        filled = generate_filled_document(buf.getvalue(), "t.xlsx", {"{{amount}}": "1+1"})
        cell = openpyxl.load_workbook(io.BytesIO(filled)).active["A1"]
        assert cell.data_type == "s"
        assert cell.value == "'=1+1"


class TestGeneratorValueSanitizing:
    def test_control_characters_do_not_crash_generation(self):
        tpl = _docx_bytes(lambda d: d.add_paragraph("Value: {{v}}"))
        filled = generate_filled_document(tpl, "t.docx", {"{{v}}": "Line\x0bBreak\x00!"})
        text = "\n".join(p.text for p in Document(io.BytesIO(filled)).paragraphs)
        assert "LineBreak!" in text

    def test_single_asterisk_arithmetic_preserved(self):
        assert _clean_field_value("5 * 3 * 2") == "5 * 3 * 2"
        assert _clean_field_value("2 * x * 3") == "2 * x * 3"

    def test_markdown_bold_stripped(self):
        assert _clean_field_value("**bold** text") == "bold text"


class TestGeneratorRunFormatting:
    def test_mixed_run_formatting_is_preserved(self):
        def build(doc):
            p = doc.add_paragraph()
            r1 = p.add_run("Dear ")
            r1.bold = True
            p.add_run("{{name}}")
            r3 = p.add_run(" please pay.")
            r3.italic = True

        tpl = _docx_bytes(build)
        filled = generate_filled_document(tpl, "t.docx", {"{{name}}": "Alice"})
        para = Document(io.BytesIO(filled)).paragraphs[0]
        assert para.text == "Dear Alice please pay."
        runs = para.runs
        assert runs[0].bold is True
        assert runs[0].text == "Dear "
        assert any(r.italic is True and "please pay" in r.text for r in runs)


class TestGeneratorFooterAndNestedTables:
    def test_footer_table_placeholder_is_filled(self):
        from docx.shared import Inches

        def build(doc):
            table = doc.sections[0].footer.add_table(rows=1, cols=1, width=Inches(3))
            table.rows[0].cells[0].text = "{{ftr}}"

        tpl = _docx_bytes(build)
        parsed = parse_template_bytes(tpl, filename="t.docx")
        assert "ftr" in parsed.field_names
        filled = generate_filled_document(tpl, "t.docx", {"{{ftr}}": "Filled"})
        out = Document(io.BytesIO(filled))
        assert out.sections[0].footer.tables[0].rows[0].cells[0].text == "Filled"

    def test_nested_table_placeholder_is_filled(self):
        def build(doc):
            outer = doc.add_table(rows=1, cols=1)
            inner = outer.rows[0].cells[0].add_table(rows=1, cols=1)
            inner.rows[0].cells[0].text = "{{nested}}"

        tpl = _docx_bytes(build)
        parsed = parse_template_bytes(tpl, filename="t.docx")
        assert "nested" in parsed.field_names
        filled = generate_filled_document(tpl, "t.docx", {"{{nested}}": "Deep"})
        out = Document(io.BytesIO(filled))
        inner = out.tables[0].rows[0].cells[0].tables[0]
        assert inner.rows[0].cells[0].text == "Deep"


# ---------------------------------------------------------------------------
# PII masker
# ---------------------------------------------------------------------------

class TestMaskerLeaks:
    def test_full_prefixed_phone_number_is_masked(self):
        masker = SelectivePIIMasker()
        masked, mapping = masker.mask_text("Telepon: (031) 567-8901")
        assert "8901" not in masked
        assert "(031) 567-8901" not in masked
        assert masker.unmask(masked, mapping) == "Telepon: (031) 567-8901"

    def test_wa_number_with_country_code_fully_masked(self):
        masker = SelectivePIIMasker()
        masked, _ = masker.mask_text("wa +62 812-3456-7890")
        assert "7890" not in masked

    def test_unmask_more_than_ten_tokens_of_one_category(self):
        masker = SelectivePIIMasker()
        mapping = {f"[TOKEN_PHONE_{i}]": f"0812345678{i:02d}" for i in range(1, 12)}
        out = masker.unmask("[TOKEN_PHONE_10] [TOKEN_PHONE_1]", mapping)
        assert out == "081234567810 081234567801"

    def test_unmask_handles_backslashes_in_value(self):
        masker = SelectivePIIMasker()
        out = masker.unmask("[TOKEN_EMAIL_1]", {"[TOKEN_EMAIL_1]": r"C:\Users\1\data"})
        assert out == r"C:\Users\1\data"

    def test_unmask_restores_lowercase_bare_token(self):
        masker = SelectivePIIMasker()
        out = masker.unmask("token_npwp_1", {"[TOKEN_NPWP_1]": "01.234.567.8-615.000"})
        assert out == "01.234.567.8-615.000"


# ---------------------------------------------------------------------------
# Mapper
# ---------------------------------------------------------------------------

class TestMapperAliases:
    def test_date_fields_do_not_cross_fill(self):
        tpl = _docx_bytes(lambda d: d.add_paragraph("{{due_date}}"))
        parsed = parse_template_bytes(tpl, filename="t.docx")
        result = map_fields({"invoice_date": "2026-01-01"}, parsed)
        assert "due_date" in result.unmapped
        assert result.mapped == {}

    def test_amount_does_not_cross_fill_total(self):
        tpl = _docx_bytes(lambda d: d.add_paragraph("{{amount}}"))
        parsed = parse_template_bytes(tpl, filename="t.docx")
        result = map_fields({"total": "999"}, parsed)
        assert "amount" in result.unmapped

    def test_short_name_fuzzy_match_rejected(self):
        tpl = _docx_bytes(lambda d: d.add_paragraph("{{date}}"))
        parsed = parse_template_bytes(tpl, filename="t.docx")
        result = map_fields({"rate": "7.5%"}, parsed)
        assert "date" in result.unmapped


# ---------------------------------------------------------------------------
# Chunker
# ---------------------------------------------------------------------------

class TestChunkerHeaderAndOverlap:
    def test_chunk_header_reflects_its_own_section(self):
        from app.services.extraction import ExtractedDocument, PageContent, PdfMetadata

        intro = "Introduction paragraph. " * 40
        methods = "Methods paragraph. " * 40
        text = f"# Introduction\n{intro}\n\n## Methods\n{methods}"
        doc = ExtractedDocument(
            filename="t.pdf",
            metadata=PdfMetadata(page_count=1, has_text=True),
            pages=[PageContent(page_number=1, text=text, width=595, height=842, blocks=[], tables=[], char_count=len(text), has_text=True)],
            tables=[],
            full_text=text,
        )
        chunks = chunk_document(doc, chunk_size=100, chunk_overlap=10)
        # A chunk that STARTS in the introduction must be labelled Introduction;
        # one that starts in the methods section must be labelled Methods.
        for c in chunks:
            if c.text.startswith("# Introduction") or c.text.startswith("Introduction paragraph"):
                assert c.header == "Introduction", c.text[:60]
            if c.text.startswith("## Methods") or c.text.startswith("Methods paragraph"):
                assert c.header == "Methods", c.text[:60]

    def test_overlap_does_not_compound_from_previous_overlap(self):
        chunks = ["a" * 400, "b" * 400, "c" * 400]
        overlapped = _apply_overlap(chunks, chunk_size=100, chunk_overlap=20)
        # Third chunk's prepended overlap must come from the ORIGINAL second
        # chunk (all 'b'), never from the already-overlapped second chunk.
        assert "a" not in overlapped[2]


# ---------------------------------------------------------------------------
# Embedder / key pool
# ---------------------------------------------------------------------------

class TestEmbedderAndKeyPool:
    def test_fake_embedding_respects_dims(self):
        assert len(_fake_embedding("x", 8)) == 8
        assert len(Embedder(use_fake=True, dims=16).embed_texts_sync(["hello"])[0]) == 16

    def test_explicit_key_does_not_clobber_global_pool(self):
        before = get_key_pool().get_all_keys()
        from app.services.generation.extractor import GeminiExtractor

        ext = GeminiExtractor(api_key="k1, k2", use_fake=True)
        assert ext.key_pool.get_all_keys() == ["k1", "k2"]
        assert get_key_pool().get_all_keys() == before

    def test_standalone_pool_is_independent(self):
        pool = GeminiKeyPool("only-key")
        assert pool.get_all_keys() == ["only-key"]


# ---------------------------------------------------------------------------
# Vector store / table extractor
# ---------------------------------------------------------------------------

class TestVectorStoreParity:
    def test_dim_mismatch_skipped_in_both_search_paths(self):
        store = InMemoryVectorStore()
        store.add_sync([StoredChunk(chunk_id="c1", text="t", embedding=[0.1] * 768, page_number=1)])
        query = [0.1] * 3072  # wrong dims
        import asyncio

        assert asyncio.run(store.search(query, top_k=5)) == []
        assert store.search_sync(query, top_k=5) == []


class TestTableExtractorNoneRow:
    def test_none_row_does_not_raise(self):
        table = _table_to_extracted([None, ["Alice", "100"]], page_number=1, table_index=0)
        assert table is not None
        assert table.headers == ["Alice", "100"]

    def test_all_none_rows_returns_none(self):
        assert _table_to_extracted([None, None], page_number=1, table_index=0) is None
