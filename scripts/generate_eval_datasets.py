"""Generate synthetic evaluation datasets for TemplaFill (Task 4.1).

Creates 5 document pairs across 5 domains with ground_truth.json.
Source PDFs and templates are generated synthetically (no PII) so they are safe
for public repo. PDFs/templates are written to eval/datasets/ but are gitignored
per .gitignore; run this script to regenerate them before eval.

Usage:
    python scripts/generate_eval_datasets.py
"""

from __future__ import annotations

import io
import json
import pathlib

import fitz  # PyMuPDF
from docx import Document
import openpyxl
from pptx import Presentation
from pptx.util import Inches

ROOT = pathlib.Path(__file__).parent.parent
DATASETS_ROOT = ROOT / "eval" / "datasets"

# ---------------------------------------------------------------------------
# Dataset specs
# ---------------------------------------------------------------------------

DATASETS = [
    {
        "domain": "hr",
        "dataset_id": "hr_cv_01",
        "source_file": "cv_sample_01.pdf",
        "template_file": "offer_letter_template.docx",
        "template_format": "docx",
        "fields": [
            {"field_name": "full_name", "placeholder": "{{full_name}}", "expected_value": "John Doe", "is_present": True},
            {"field_name": "email", "placeholder": "{{email}}", "expected_value": "john.doe@example.com", "is_present": True},
            {"field_name": "phone_number", "placeholder": "{{phone_number}}", "expected_value": "0812-3456-7890", "is_present": True},
            {"field_name": "position", "placeholder": "{{position}}", "expected_value": "Software Engineer", "is_present": True},
            {"field_name": "skills", "placeholder": "{{skills}}", "expected_value": "Python, FastAPI, React", "is_present": True},
            {"field_name": "address", "placeholder": "{{address}}", "expected_value": None, "is_present": False},
        ],
        "source_lines": [
            "Curriculum Vitae",
            "Full Name: John Doe",
            "Email: john.doe@example.com",
            "Phone Number: 0812-3456-7890",
            "Position: Software Engineer",
            "Skills: Python, FastAPI, React",
            "Experience: 5 years building web applications at TechCorp",
            "Education: B.Sc. Computer Science, University of Indonesia, 2020",
        ],
        "template_lines": [
            "OFFER LETTER",
            "Dear {{full_name}},",
            "We are pleased to offer you the position of {{position}} at our company.",
            "Your contact email is {{email}} and phone {{phone_number}}.",
            "Your skills in {{skills}} are highly valued.",
            "Address on file: {{address}}",
            "Sincerely, HR Department",
        ],
    },
    {
        "domain": "finance",
        "dataset_id": "finance_01",
        "source_file": "financial_report_01.pdf",
        "template_file": "quarterly_summary.xlsx",
        "template_format": "xlsx",
        "fields": [
            {"field_name": "invoice_number", "placeholder": "{{invoice_number}}", "expected_value": "INV-2026-001", "is_present": True},
            {"field_name": "invoice_date", "placeholder": "{{invoice_date}}", "expected_value": "2026-03-15", "is_present": True},
            {"field_name": "total_amount", "placeholder": "{{total_amount}}", "expected_value": "$12,300", "is_present": True},
            {"field_name": "customer_name", "placeholder": "{{customer_name}}", "expected_value": "Acme Corp", "is_present": True},
            {"field_name": "due_date", "placeholder": "{{due_date}}", "expected_value": None, "is_present": False},
        ],
        "source_lines": [
            "Financial Report Q1 2026",
            "Invoice Number: INV-2026-001",
            "Invoice Date: 2026-03-15",
            "Customer Name: Acme Corp",
            "Total Amount: $12,300",
            "Items: Widget A x10 $5,000, Widget B x5 $7,300",
            "Payment Terms: Net 30",
        ],
        "template_cells": {
            "A1": "Quarterly Summary - {{invoice_number}}",
            "A2": "Date: {{invoice_date}}",
            "A3": "Customer: {{customer_name}}",
            "A4": "Total: {{total_amount}}",
            "A5": "Due: {{due_date}}",
        },
    },
    {
        "domain": "education",
        "dataset_id": "education_01",
        "source_file": "transcript_01.pdf",
        "template_file": "certificate_template.docx",
        "template_format": "docx",
        "fields": [
            {"field_name": "student_name", "placeholder": "{{student_name}}", "expected_value": "Alice Smith", "is_present": True},
            {"field_name": "gpa", "placeholder": "{{gpa}}", "expected_value": "3.85", "is_present": True},
            {"field_name": "graduation_date", "placeholder": "{{graduation_date}}", "expected_value": "2026-06-20", "is_present": True},
            {"field_name": "major", "placeholder": "{{major}}", "expected_value": "Computer Science", "is_present": True},
            {"field_name": "honors", "placeholder": "{{honors}}", "expected_value": None, "is_present": False},
        ],
        "source_lines": [
            "Academic Transcript",
            "Student Name: Alice Smith",
            "Major: Computer Science",
            "GPA: 3.85",
            "Graduation Date: 2026-06-20",
            "Courses: Algorithms A, Database A, Thesis A",
            "University: Institut Teknologi Bandung",
        ],
        "template_lines": [
            "CERTIFICATE OF GRADUATION",
            "This certifies that {{student_name}}",
            "has completed major in {{major}}",
            "with GPA {{gpa}} on {{graduation_date}}",
            "Honors: {{honors}}",
        ],
    },
    {
        "domain": "legal",
        "dataset_id": "legal_01",
        "source_file": "contract_01.pdf",
        "template_file": "summary_template.docx",
        "template_format": "docx",
        "fields": [
            {"field_name": "contract_number", "placeholder": "{{contract_number}}", "expected_value": "CTR-2026-889", "is_present": True},
            {"field_name": "signing_date", "placeholder": "{{signing_date}}", "expected_value": "2026-08-15", "is_present": True},
            {"field_name": "party_name", "placeholder": "{{party_name}}", "expected_value": "PT Maju Jaya", "is_present": True},
            {"field_name": "contract_value", "placeholder": "{{contract_value}}", "expected_value": "$50,000", "is_present": True},
            {"field_name": "witness_name", "placeholder": "{{witness_name}}", "expected_value": None, "is_present": False},
        ],
        "source_lines": [
            "Contract Agreement",
            "Contract Number: CTR-2026-889",
            "Signing Date: 2026-08-15",
            "Party Name: PT Maju Jaya",
            "Contract Value: $50,000",
            "Terms: This agreement is valid for 12 months and governed by Indonesian law.",
            "Signature block follows on last page.",
        ],
        "template_lines": [
            "CONTRACT SUMMARY",
            "Contract {{contract_number}} signed on {{signing_date}}",
            "between PT Maju Jaya and counterparty.",
            "Party: {{party_name}}",
            "Value: {{contract_value}}",
            "Witness: {{witness_name}}",
        ],
    },
    {
        "domain": "general",
        "dataset_id": "general_01",
        "source_file": "report_01.pdf",
        "template_file": "brief_template.docx",
        "template_format": "docx",
        "fields": [
            {"field_name": "company_name", "placeholder": "{{company_name}}", "expected_value": "IndoTech Solutions", "is_present": True},
            {"field_name": "report_date", "placeholder": "{{report_date}}", "expected_value": "2026-09-01", "is_present": True},
            {"field_name": "total_employees", "placeholder": "{{total_employees}}", "expected_value": "150", "is_present": True},
            {"field_name": "ceo_name", "placeholder": "{{ceo_name}}", "expected_value": "Budi Santoso", "is_present": True},
            {"field_name": "revenue", "placeholder": "{{revenue}}", "expected_value": None, "is_present": False},
        ],
        "source_lines": [
            "Annual Report 2026",
            "Company Name: IndoTech Solutions",
            "Report Date: 2026-09-01",
            "CEO Name: Budi Santoso",
            "Total Employees: 150",
            "Overview: Company achieved growth in Q1-Q3.",
            "Headquarters: Jakarta, Indonesia",
        ],
        "template_lines": [
            "EXECUTIVE BRIEF",
            "Company: {{company_name}}",
            "Date: {{report_date}}",
            "CEO: {{ceo_name}}",
            "Employees: {{total_employees}}",
            "Revenue: {{revenue}}",
        ],
    },
]

# Add pptx domain for variety? We already have 5; add extra for coverage
# General already covers 5 domains, meets MVP 5 datasets across 5 domains


def create_pdf(path: pathlib.Path, lines: list[str], title: str = ""):
    doc = fitz.open()
    page = doc.new_page()
    rect = fitz.Rect(50, 50, 550, 800)
    text = "\n".join(lines)
    page.insert_textbox(rect, text, fontsize=11, fontname="helv")
    if title:
        doc.set_metadata({"title": title})
    buf = io.BytesIO()
    doc.save(buf)
    doc.close()
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(buf.getvalue())
    print(f"Created PDF: {path} ({len(lines)} lines)")


def create_docx_template(path: pathlib.Path, lines: list[str]):
    doc = Document()
    for line in lines:
        doc.add_paragraph(line)
    path.parent.mkdir(parents=True, exist_ok=True)
    buf = io.BytesIO()
    doc.save(buf)
    path.write_bytes(buf.getvalue())
    print(f"Created DOCX: {path} ({len(lines)} lines)")


def create_xlsx_template(path: pathlib.Path, cells: dict):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Summary"
    for coord, val in cells.items():
        ws[coord] = val
    path.parent.mkdir(parents=True, exist_ok=True)
    buf = io.BytesIO()
    wb.save(buf)
    path.write_bytes(buf.getvalue())
    print(f"Created XLSX: {path} ({len(cells)} cells)")


def create_pptx_template(path: pathlib.Path, lines: list[str]):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    left, top, width, height = Inches(0.5), Inches(0.5), Inches(9), Inches(5)
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    for line in lines:
        p = tf.add_paragraph()
        p.text = line
    path.parent.mkdir(parents=True, exist_ok=True)
    buf = io.BytesIO()
    prs.save(buf)
    path.write_bytes(buf.getvalue())
    print(f"Created PPTX: {path}")


def main():
    for ds in DATASETS:
        domain_dir = DATASETS_ROOT / ds["domain"]
        domain_dir.mkdir(parents=True, exist_ok=True)
        # Source PDF
        pdf_path = domain_dir / ds["source_file"]
        create_pdf(pdf_path, ds["source_lines"], title=ds["dataset_id"])
        # Template
        tmpl_path = domain_dir / ds["template_file"]
        fmt = ds["template_format"]
        if fmt == "docx":
            create_docx_template(tmpl_path, ds["template_lines"])
        elif fmt == "xlsx":
            create_xlsx_template(tmpl_path, ds["template_cells"])
        elif fmt == "pptx":
            create_pptx_template(tmpl_path, ds["template_lines"])
        else:
            raise ValueError(fmt)
        # Ground truth JSON
        gt_path = domain_dir / "ground_truth_01.json"
        gt = {
            "dataset_id": ds["dataset_id"],
            "source_file": ds["source_file"],
            "template_file": ds["template_file"],
            "fields": [
                {
                    "field_name": f["field_name"],
                    "placeholder": f["placeholder"],
                    "expected_value": f["expected_value"],
                    "source_page": 1 if f["is_present"] else None,
                    "is_present_in_source": f["is_present"],
                }
                for f in ds["fields"]
            ],
        }
        gt_path.write_text(json.dumps(gt, indent=2))
        print(f"Created ground truth: {gt_path}")
        # README
        readme = domain_dir / "README.md"
        readme.write_text(
            f"# {ds['dataset_id']}\n\nDomain: {ds['domain']}\nFields: {len(ds['fields'])}\n"
            f"Source: {ds['source_file']}, Template: {ds['template_file']}\n"
            f"Generated synthetically for eval — no PII.\n"
        )
    print(f"\nDone: {len(DATASETS)} datasets created under {DATASETS_ROOT}")


if __name__ == "__main__":
    main()
